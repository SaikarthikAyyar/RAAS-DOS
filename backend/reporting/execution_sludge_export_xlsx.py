# ====================================
# EXECUTION DAILY SLUDGE OUTPUT - STYLED .XLSX EXPORT
# One sheet per day (mirroring the day-per-page layout of the real
# client worksheets this whole feature's formulas were verified
# against - "Daily Sludge Calculation Day 7-Day 10.pdf" /
# "...8 May.pdf"), plus one Summary sheet up front. Same shared visual
# design as every other server-side export in this app
# (backend/reporting/xlsx_style.py) - not a second, hand-rolled style.
# ====================================

from io import BytesIO

from openpyxl import Workbook

from backend.reporting.xlsx_style import (
    title_row,
    subtitle_row,
    section_band,
    field_row,
    table_header_row,
    table_data_row,
    sheet_name as safe_sheet_name
)


METHOD_LABELS = {
    "FLOW_METER": "Flow Meter Reading",
    "SETTLING": "Sample Collection (Settling)"
}


def _fmt(value, places=3):

    if value is None:
        return "-"

    return round(value, places)


def _build_summary_sheet(wb, execution, daily_logs):

    ws = wb.create_sheet(safe_sheet_name("Summary"))

    total_cols = 5

    row = 1
    title_row(ws, row, total_cols, "JANYU TECHNOLOGIES")
    row += 1

    subtitle_row(ws, row, total_cols, f"Execution #{execution.id} - Daily Sludge Output Summary")
    row += 1

    table_header_row(ws, row, ["Date", "Method", "Status", "Sludge Output (m3)", "Water Output (m3)"])
    row += 1

    total_sludge = 0
    total_water = 0

    for log in daily_logs:

        status = (
            "Invalid - see sheet" if log.invalid_reason
            else "Complete" if log.sludge_output_m3 is not None
            else "Pending"
        )

        table_data_row(ws, row, [
            log.log_date.isoformat() if log.log_date else "-",
            METHOD_LABELS.get(log.method, log.method),
            status,
            _fmt(log.sludge_output_m3),
            _fmt(log.water_output_m3)
        ])
        row += 1

        total_sludge += log.sludge_output_m3 or 0
        total_water += log.water_output_m3 or 0

    table_data_row(ws, row, ["", "", "TOTAL", round(total_sludge, 3), round(total_water, 3)])

    for i, width in enumerate([14, 26, 14, 20, 20], start=1):
        ws.column_dimensions[ws.cell(row=4, column=i).column_letter].width = width


def _build_day_sheet(wb, execution, log):

    ws = wb.create_sheet(safe_sheet_name(log.log_date.isoformat() if log.log_date else f"Day {log.id}"))

    total_cols = 3 if log.method == "FLOW_METER" else 4

    row = 1
    title_row(ws, row, total_cols, "JANYU TECHNOLOGIES")
    row += 1

    subtitle_row(ws, row, total_cols, f"Execution #{execution.id} - {log.log_date} - {METHOD_LABELS.get(log.method, log.method)}")
    row += 1

    field_row(ws, row, "Start TF (m3)", [_fmt(log.start_tf)])
    row += 1

    field_row(ws, row, "End TF (m3)", [_fmt(log.end_tf)])
    row += 1

    if log.method == "FLOW_METER":
        field_row(ws, row, "Total Sludge Pump Time (min)", [_fmt(log.total_sludge_pump_minutes, 1)])
        row += 1

    row += 1

    section_band(ws, row, total_cols, "READINGS")
    row += 1

    if log.method == "FLOW_METER":
        headers = ["Time", "TF (m3)", "FR (m3/hr)"]
    else:
        # Flask volume is captured per reading, not once for the whole
        # day - different flask sizes may genuinely be used sample to
        # sample.
        headers = ["Time", "TF (m3)", "Settled Sludge (ml)", "Flask Volume (ml)"]

    table_header_row(ws, row, headers)
    row += 1

    for reading in log.readings:

        if log.method == "FLOW_METER":
            values = [
                reading.recorded_at.strftime("%Y-%m-%d %H:%M") if reading.recorded_at else "-",
                _fmt(reading.tf_reading),
                _fmt(reading.fr_reading, 2)
            ]
        else:
            values = [
                reading.recorded_at.strftime("%Y-%m-%d %H:%M") if reading.recorded_at else "-",
                _fmt(reading.tf_reading),
                _fmt(reading.settled_sludge_volume_ml, 0),
                _fmt(reading.flask_volume_ml, 0)
            ]

        table_data_row(ws, row, values)
        row += 1

    row += 1

    section_band(ws, row, total_cols, "COMPUTED RESULTS")
    row += 1

    if log.invalid_reason:

        field_row(ws, row, "Not Computed - Check Data", [log.invalid_reason], wrap=True)
        row += 1

        ws.column_dimensions["A"].width = 30
        ws.column_dimensions["B"].width = 22
        ws.column_dimensions["C"].width = 22

        if log.method == "SETTLING":
            ws.column_dimensions["D"].width = 22

        return

    if log.method == "FLOW_METER":

        field_row(ws, row, "Avg FR (m3/hr)", [_fmt(log.avg_fr, 2)])
        row += 1

        field_row(ws, row, "FR per Minute (m3/min)", [_fmt(log.fr_per_minute, 4)])
        row += 1

        field_row(ws, row, "Estimated Volume (m3)", [_fmt(log.total_sludge_pumping_estimate_m3, 2)])
        row += 1

    field_row(ws, row, "Total TF - ground truth (m3)", [_fmt(log.total_tf_m3)])
    row += 1

    field_row(ws, row, "% Sludge / % Water", [
        f"{_fmt(log.pct_sludge, 1)}% / {_fmt(log.pct_water, 1)}%"
    ])
    row += 1

    field_row(ws, row, "Sludge Output (m3)", [_fmt(log.sludge_output_m3)])
    row += 1

    field_row(ws, row, "Water Output (m3)", [_fmt(log.water_output_m3)])
    row += 1

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 22

    if log.method == "SETTLING":
        ws.column_dimensions["D"].width = 22


def build_execution_sludge_workbook_bytes(execution, daily_logs):

    wb = Workbook()
    wb.remove(wb.active)

    _build_summary_sheet(wb, execution, daily_logs)

    for log in daily_logs:
        _build_day_sheet(wb, execution, log)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return buffer
